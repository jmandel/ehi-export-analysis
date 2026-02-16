"""
Parse b10_data_dictionary.xlsx and produce entity-inventory-full.json and entity-inventory-summary.json.

Input:  ../downloads/b10_data_dictionary.xlsx
Output: entity-inventory-full.json   — complete extraction of all entities and fields
        entity-inventory-summary.json — aggregate statistics and category breakdowns
"""

import json
import openpyxl
from collections import defaultdict, OrderedDict

INPUT = "../downloads/b10_data_dictionary.xlsx"

wb = openpyxl.load_workbook(INPUT, read_only=True, data_only=True)
print(f"Sheets: {wb.sheetnames}")

entities = []

# --- Sheet 1: iKnowMed ---
ws = wb["iKnowMed"]
rows = list(ws.iter_rows(min_row=1, values_only=False))
headers = [c.value for c in rows[0]]
print(f"iKnowMed headers: {headers}")

by_table = OrderedDict()
for row in rows[1:]:
    vals = {headers[i]: row[i].value for i in range(len(headers))}
    tname = str(vals.get("TABLE_NAME") or "").strip()
    if not tname:
        continue
    if tname not in by_table:
        by_table[tname] = {
            "table_desc": str(vals.get("TABLE_DESC") or "").strip() or None,
            "columns": []
        }
    col_desc = str(vals.get("COLUMN_DESC") or "").strip() or None
    by_table[tname]["columns"].append({
        "column_name": str(vals.get("COLUMN_NAME") or "").strip(),
        "data_type": str(vals.get("DATA_TYPE") or "").strip(),
        "data_length": vals.get("DATA_LENGTH"),
        "description": col_desc,
    })

for tname, info in by_table.items():
    entities.append({
        "entity_name": tname,
        "source_sheet": "iKnowMed",
        "entity_type": "table",
        "entity_description": info["table_desc"],
        "fields": info["columns"],
    })

print(f"  iKnowMed: {len(by_table)} tables, {sum(len(t['columns']) for t in by_table.values())} columns")

# --- Sheet 2: VBC ---
ws = wb["VBC"]
rows = list(ws.iter_rows(min_row=1, values_only=False))
headers_vbc = [c.value for c in rows[0]]
print(f"VBC headers: {headers_vbc}")

by_table_vbc = OrderedDict()
for row in rows[1:]:
    vals = {headers_vbc[i]: row[i].value for i in range(len(headers_vbc))}
    tname = str(vals.get("table_name") or "").strip()
    if not tname:
        continue
    if tname not in by_table_vbc:
        by_table_vbc[tname] = []
    # Check for description columns
    desc = None
    for key in ["column_desc", "COLUMN_DESC", "description", "Description"]:
        if vals.get(key):
            desc = str(vals[key]).strip()
            break
    by_table_vbc[tname].append({
        "column_name": str(vals.get("column_name") or "").strip(),
        "data_type": str(vals.get("data_type") or "").strip(),
        "data_length": None,
        "description": desc,
    })

for tname, cols in by_table_vbc.items():
    entities.append({
        "entity_name": tname,
        "source_sheet": "VBC",
        "entity_type": "table",
        "entity_description": None,
        "fields": cols,
    })

print(f"  VBC: {len(by_table_vbc)} tables, {sum(len(c) for c in by_table_vbc.values())} columns")

# --- Sheet 3: Patient History ---
ws = wb["Patient History"]
rows = list(ws.iter_rows(min_row=1, values_only=False))
headers_ph = [c.value for c in rows[0]]
print(f"Patient History headers: {headers_ph}")

by_table_ph = OrderedDict()
for row in rows[1:]:
    vals = {headers_ph[i]: row[i].value for i in range(len(headers_ph))}
    tname = str(vals.get("TABLE_NAME") or "").strip()
    if not tname:
        continue
    if tname not in by_table_ph:
        by_table_ph[tname] = []
    desc = None
    for key in ["COLUMN_DESC", "column_desc", "Description"]:
        if vals.get(key):
            desc = str(vals[key]).strip()
            break
    by_table_ph[tname].append({
        "column_name": str(vals.get("COLUMN_NAME") or "").strip(),
        "data_type": str(vals.get("DATA_TYPE") or "").strip(),
        "data_length": vals.get("DATA_LENGTH"),
        "description": desc,
    })

for tname, cols in by_table_ph.items():
    entities.append({
        "entity_name": tname,
        "source_sheet": "Patient History",
        "entity_type": "table",
        "entity_description": None,
        "fields": cols,
    })

print(f"  Patient History: {len(by_table_ph)} tables, {sum(len(c) for c in by_table_ph.values())} columns")

# --- Sheet 4: Ontada Health ---
ws = wb["Ontada Health"]
rows = list(ws.iter_rows(min_row=1, values_only=False))
headers_oh = [c.value for c in rows[0]]
print(f"Ontada Health headers: {headers_oh}")

by_coll = OrderedDict()
for row in rows[1:]:
    vals = {headers_oh[i]: row[i].value for i in range(len(headers_oh))}
    cname = str(vals.get("Collection") or "").strip()
    if not cname:
        continue
    if cname not in by_coll:
        by_coll[cname] = []
    desc = str(vals.get("Description") or "").strip() or None
    by_coll[cname].append({
        "column_name": str(vals.get("Field Name") or "").strip(),
        "data_type": str(vals.get("Data Type") or "").strip(),
        "data_length": None,
        "description": desc,
    })

for cname, fields in by_coll.items():
    entities.append({
        "entity_name": cname,
        "source_sheet": "Ontada Health",
        "entity_type": "collection",
        "entity_description": None,
        "fields": fields,
    })

print(f"  Ontada Health: {len(by_coll)} collections, {sum(len(f) for f in by_coll.values())} fields")

wb.close()

# --- Compute statistics ---
total_entities = len(entities)
total_fields = sum(len(e["fields"]) for e in entities)
fields_with_desc = sum(
    1 for e in entities for f in e["fields"]
    if f["description"] and f["description"].strip()
)
fields_with_types = sum(
    1 for e in entities for f in e["fields"]
    if f["data_type"] and f["data_type"].strip()
)
entities_with_desc = sum(1 for e in entities if e["entity_description"])

# By sheet
sheet_stats = defaultdict(lambda: {"entities": 0, "fields": 0, "fields_with_desc": 0})
for e in entities:
    s = e["source_sheet"]
    sheet_stats[s]["entities"] += 1
    sheet_stats[s]["fields"] += len(e["fields"])
    sheet_stats[s]["fields_with_desc"] += sum(
        1 for f in e["fields"] if f["description"] and f["description"].strip()
    )

# --- Write full inventory ---
full_output = {
    "extraction_date": "2026-02-16",
    "source_file": "b10_data_dictionary.xlsx",
    "total_entities": total_entities,
    "total_fields": total_fields,
    "entities": entities,
}

with open("entity-inventory-full.json", "w") as f:
    json.dump(full_output, f, indent=2)
print(f"\nWrote entity-inventory-full.json ({total_entities} entities, {total_fields} fields)")

# --- Write summary ---
# Top 20 largest entities
largest = sorted(entities, key=lambda e: len(e["fields"]), reverse=True)[:20]

summary = {
    "extraction_date": "2026-02-16",
    "source_file": "b10_data_dictionary.xlsx",
    "total_entities": total_entities,
    "total_fields": total_fields,
    "fields_with_descriptions": fields_with_desc,
    "fields_with_descriptions_pct": round(fields_with_desc / total_fields * 100, 1) if total_fields else 0,
    "fields_with_types": fields_with_types,
    "entities_with_descriptions": entities_with_desc,
    "by_sheet": {k: dict(v) for k, v in sheet_stats.items()},
    "top_20_largest_entities": [
        {
            "entity_name": e["entity_name"],
            "source_sheet": e["source_sheet"],
            "field_count": len(e["fields"]),
            "has_description": e["entity_description"] is not None,
            "fields_with_desc": sum(1 for f in e["fields"] if f["description"] and f["description"].strip()),
        }
        for e in largest
    ],
    "all_entities": [
        {
            "entity_name": e["entity_name"],
            "source_sheet": e["source_sheet"],
            "field_count": len(e["fields"]),
            "fields_with_desc": sum(1 for f in e["fields"] if f["description"] and f["description"].strip()),
        }
        for e in entities
    ],
}

with open("entity-inventory-summary.json", "w") as f:
    json.dump(summary, f, indent=2)
print(f"Wrote entity-inventory-summary.json")
print(f"\nStats: {total_entities} entities, {total_fields} fields, {fields_with_desc} with descriptions ({round(fields_with_desc/total_fields*100,1)}%), {entities_with_desc} entities with descriptions")
print(f"By sheet: {json.dumps({k: dict(v) for k,v in sheet_stats.items()}, indent=2)}")
