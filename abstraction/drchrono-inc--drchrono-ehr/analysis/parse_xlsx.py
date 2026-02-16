#!/usr/bin/env python3
"""Parse drchrono EHI Export XLSX data dictionary v1.8 into structured JSON."""

import json
import openpyxl
from collections import defaultdict

XLSX_PATH = "../downloads/drchrono-ehi-export-documentation-v18.xlsx"

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

results = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue
    
    # Find header row
    headers = [str(c).strip() if c else "" for c in rows[0]]
    print(f"\nSheet: {sheet_name}")
    print(f"  Headers: {headers}")
    print(f"  Total rows (incl header): {len(rows)}")
    
    # Parse data rows
    entities = defaultdict(list)
    for row in rows[1:]:
        vals = list(row)
        if len(vals) < 2:
            continue
        export_name = str(vals[0]).strip() if vals[0] else ""
        field_name = str(vals[1]).strip() if vals[1] else ""
        description = str(vals[2]).strip() if len(vals) > 2 and vals[2] else ""
        
        if not export_name and not field_name:
            continue
            
        entities[export_name].append({
            "field": field_name,
            "description": description
        })
    
    entity_list = []
    for ename, fields in entities.items():
        entity_list.append({
            "entity_name": ename,
            "field_count": len(fields),
            "fields_with_description": sum(1 for f in fields if f["description"]),
            "fields": fields
        })
    
    results[sheet_name] = {
        "entity_count": len(entity_list),
        "total_fields": sum(e["field_count"] for e in entity_list),
        "total_fields_with_description": sum(e["fields_with_description"] for e in entity_list),
        "entities": entity_list
    }
    
    print(f"  Entities: {len(entity_list)}")
    print(f"  Total fields: {sum(e['field_count'] for e in entity_list)}")
    print(f"  Fields with descriptions: {sum(e['fields_with_description'] for e in entity_list)}")

wb.close()

# Save full inventory
with open("entity-inventory-full.json", "w") as f:
    json.dump(results, f, indent=2)

# Build summary
summary = {}
for sheet_name, data in results.items():
    entity_summary = []
    for e in data["entities"]:
        entity_summary.append({
            "entity_name": e["entity_name"],
            "field_count": e["field_count"],
            "fields_with_description": e["fields_with_description"],
            "pct_described": round(100 * e["fields_with_description"] / e["field_count"], 1) if e["field_count"] > 0 else 0
        })
    entity_summary.sort(key=lambda x: -x["field_count"])
    summary[sheet_name] = {
        "entity_count": data["entity_count"],
        "total_fields": data["total_fields"],
        "total_fields_with_description": data["total_fields_with_description"],
        "pct_described": round(100 * data["total_fields_with_description"] / data["total_fields"], 1) if data["total_fields"] > 0 else 0,
        "entities": entity_summary
    }

with open("entity-inventory-summary.json", "w") as f:
    json.dump(summary, f, indent=2)

print("\n\nDone. Saved entity-inventory-full.json and entity-inventory-summary.json")
