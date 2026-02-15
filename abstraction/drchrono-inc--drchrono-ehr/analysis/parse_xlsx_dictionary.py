#!/usr/bin/env python3
"""Parse drchrono EHI Export XLSX data dictionary v1.8 and produce summary statistics."""
import json
import openpyxl

XLSX_PATH = "../../../results/drchrono-inc--drchrono-ehr/downloads/drchrono-ehi-export-documentation-v18.xlsx"

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

results = {}
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]

    # Build entity inventory
    entities = {}
    for row in data:
        export_name = row[0]
        field = row[1]
        description = row[2] if len(row) > 2 else None
        if export_name is None:
            continue
        export_name = str(export_name).strip()
        if export_name not in entities:
            entities[export_name] = {"fields": [], "described_count": 0, "total_fields": 0}
        entities[export_name]["fields"].append({
            "field": str(field).strip() if field else "",
            "description": str(description).strip() if description else "",
        })
        entities[export_name]["total_fields"] += 1
        if description and str(description).strip():
            entities[export_name]["described_count"] += 1

    total_fields = sum(e["total_fields"] for e in entities.values())
    total_described = sum(e["described_count"] for e in entities.values())

    entity_summary = []
    for name, info in sorted(entities.items()):
        entity_summary.append({
            "entity": name,
            "total_fields": info["total_fields"],
            "described_fields": info["described_count"],
            "fields": info["fields"],
        })

    results[sheet_name] = {
        "total_entities": len(entities),
        "total_fields": total_fields,
        "total_described": total_described,
        "description_pct": round(total_described / total_fields * 100, 1) if total_fields > 0 else 0,
        "entities": entity_summary,
    }

    print(f"\n=== {sheet_name} ===")
    print(f"Entities (CSV files): {len(entities)}")
    print(f"Total fields: {total_fields}")
    print(f"Fields with descriptions: {total_described} ({results[sheet_name]['description_pct']}%)")
    print(f"\nTop 20 entities by field count:")
    sorted_entities = sorted(entity_summary, key=lambda x: x["total_fields"], reverse=True)
    for e in sorted_entities[:20]:
        print(f"  {e['entity']}: {e['total_fields']} fields ({e['described_fields']} described)")

wb.close()

# Save full inventory
with open("full-entity-inventory.json", "w") as f:
    json.dump(results, f, indent=2)

print("\n\nFull inventory saved to full-entity-inventory.json")
