#!/usr/bin/env python3
"""Generate comprehensive entity inventory from XLSX v1.8, including categories and full field details.
Saves output to full-entity-inventory.json and prints summary tables for analysis.md."""
import json
import openpyxl
from collections import defaultdict

XLSX_PATH = "../../../results/drchrono-inc--drchrono-ehr/downloads/drchrono-ehi-export-documentation-v18.xlsx"

# Map every known entity to a domain category based on the glossary/reference guide
CATEGORY_MAP = {
    # Doctor Exporter
    "doctors.csv": "Doctor / Provider Info",
    "custom_vital_type.csv": "Doctor / Provider Info",
    "education_resource.csv": "Doctor / Provider Info",
    "mu_syndromic_surveillance_log.csv": "Doctor / Provider Info",
    "offices.csv": "Doctor / Provider Info",
    "patient_import_ccda_file.csv": "Doctor / Provider Info",
    "soap_note_custom_report.csv": "Doctor / Provider Info",
    "soap_note_line_item_field_type.csv": "Doctor / Provider Info",
    
    # Practice Group Exporter
    "lab_quest_cd_order_code.csv": "Practice Group",
    "lab_quest_cd_order_code_aoe.csv": "Practice Group",
    "practice_group.csv": "Practice Group",
    
    # Demographics
    "demographics.csv": "Demographics",
    "patient_occupation.csv": "Demographics",
    "uscdi_occupation_code.csv": "Demographics",
    "uscdi_industry_code.csv": "Demographics",
    "tribal_affiliations.csv": "Demographics",
    "ethnicity_subcategories.csv": "Demographics",
    "ethnicity_subcategories_1.csv": "Demographics",
    "race_subcategories.csv": "Demographics",
    "race_subcategories_1.csv": "Demographics",
    "additional_responsible_party.csv": "Demographics",
    "responsible_party_address.csv": "Demographics",
    
    # Allergies
    "allergies.csv": "Allergies",
    "allergy_snomed_code_mapping.csv": "Allergies",
    
    # Appointments
    "appointments.csv": "Encounters / Visits",
    
    # Insurance
    "auto_accident_insurance.csv": "Insurance / Coverage",
    "auto_accident_insurance_accident.csv": "Insurance / Coverage",
    "primary_hospital_insurance.csv": "Insurance / Coverage",
    "workers_comp_insurance.csv": "Insurance / Coverage",
    "insurance_authorizations.csv": "Insurance / Coverage",
    
    # Care Plans
    "care_plan.csv": "Care Plans / Goals",
    "care_plans.csv": "Care Plans / Goals",
    "care_plan_author.csv": "Care Plans / Goals",
    "care_plan_authors.csv": "Care Plans / Goals",
    "care_plan_goal_attached_codes.csv": "Care Plans / Goals",
    "care_plan_goal_problems.csv": "Care Plans / Goals",
    "care_plan_goals.csv": "Care Plans / Goals",
    "care_plan_goals": "Care Plans / Goals",
    "care_plan_intervention_attached_codes.csv": "Care Plans / Goals",
    "care_plan_goal_interventions_codes.csv": "Care Plans / Goals",
    "care_plan_interventions.csv": "Care Plans / Goals",
    "care_plan_objectives.csv": "Care Plans / Goals",
    
    # Care Team
    "care_team_member.csv": "Care Team",
    "doctor_staff_care_team_members.csv": "Care Team",
    "doctor_staff_Care_team_members.csv": "Care Team",
    "external_care_team_members.csv": "Care Team",
    "external_Care_team_members.csv": "Care Team",
    "patient_as_care_team_members.csv": "Care Team",
    
    # Claims & Billing
    "claims.csv": "Claims / Billing",
    "payments_insurance.csv": "Claims / Billing",
    "payments_patient.csv": "Claims / Billing",
    "patient_cost_estimator.csv": "Claims / Billing",
    
    # Clinical Notes
    "clinical_notes.csv": "Clinical Notes",
    "clinical_note_archives.csv": "Clinical Notes",
    "custom_clinical_note_sections.csv": "Clinical Notes",
    "note_section_comments.csv": "Clinical Notes",
    "soap_note_line_item_field_value.csv": "Clinical Notes",
    
    # Clinical Decision Support
    "clinical_decision_support_rules.csv": "Clinical Decision Support",
    "patient_specific_actions.csv": "Clinical Decision Support",
    
    # Family History
    "family_history.csv": "Family History",
    "clinical_observation.csv": "Family History",
    "clinical_observation_snomed_details.csv": "Family History",
    "person.csv": "Family History",
    "relationship.csv": "Family History",
    
    # Consent Forms
    "consent_forms.csv": "Consent Forms",
    "consent_form_assignments.csv": "Consent Forms",
    "consent_form_signatures.csv": "Consent Forms",
    "consent_form_signature_audit_logs.csv": "Consent Forms",
    
    # Communications
    "communication_log.csv": "Communications / Messages",
    "direct_message.csv": "Communications / Messages",
    "direct_message_attachment.csv": "Communications / Messages",
    "doctor_message.csv": "Communications / Messages",
    "doctor_message_log.csv": "Communications / Messages",
    "history_message.csv": "Communications / Messages",
    "outgoing_patient_message_status.csv": "Communications / Messages",
    "patient_message.csv": "Communications / Messages",
    "patient_message_attachment.csv": "Communications / Messages",
    "prescription_message.csv": "Communications / Messages",
    
    # Vitals
    "system_vitals.csv": "Vitals",
    "system_vitals_author.csv": "Vitals",
    "custom_vital_value.csv": "Vitals",
    
    # Medications
    "patient_drug.csv": "Medications / Prescriptions",
    "prescription.csv": "Medications / Prescriptions",
    "cover_my_meds_pa_request.csv": "Medications / Prescriptions",
    "pa_request_medication.csv": "Medications / Prescriptions",
    
    # Immunizations
    "patient_vaccination_record.csv": "Immunizations",
    "patientvaccinerecord_doses.csv": "Immunizations",
    "iz_patient_demographic.csv": "Immunizations",
    
    # Labs
    "lab_order.csv": "Lab Orders & Results",
    "lab_order_document.csv": "Lab Orders & Results",
    "lab_order_icd10_codes.csv": "Lab Orders & Results",
    "lab_result.csv": "Lab Orders & Results",
    "lab_result_author.csv": "Lab Orders & Results",
    "patient_lab_result_set.csv": "Lab Orders & Results",
    
    # Problems
    "problems.csv": "Problems / Diagnoses",
    
    # Social History
    "social_history.csv": "Social History",
    "social_history_author.csv": "Social History",
    
    # Functional/Mental Status
    "functional_statuses.csv": "Functional & Mental Status",
    "functional_status_authors.csv": "Functional & Mental Status",
    "mental_statuses.csv": "Functional & Mental Status",
    "mental_status_authors.csv": "Functional & Mental Status",
    
    # Devices
    "implantable_devices.csv": "Devices",
    "implantable_device_authors.csv": "Devices",
    "patient_device_orders.csv": "Devices",
    
    # Referrals
    "inbound_referrals.csv": "Referrals",
    "outbound_referrals.csv": "Referrals",
    
    # Imaging
    "patient_imaging_order.csv": "Imaging Orders",
    
    # Documents
    "uploaded_documents.csv": "Documents",
    "exported_documents": "Documents",
    
    # Procedures
    "procedures.csv": "Procedures",
    
    # Education
    "education_resource_recommendation.csv": "Patient Education",
    "mu_patient_education_log.csv": "Patient Education",
    
    # Case Reporting
    "case_report_encounters.csv": "Case Reporting",
    
    # Patient Flags
    "patient_flags.csv": "Patient Flags",
    
    # Clinical List
    "clinical_list.csv": "Imported CCDA Data",
}

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

output = {}

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    data = rows[1:]
    
    entities = {}
    for row in data:
        export_name = str(row[0]).strip() if row[0] else None
        if not export_name:
            continue
        field = str(row[1]).strip() if row[1] else ""
        desc = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        
        if export_name not in entities:
            cat = CATEGORY_MAP.get(export_name, "Uncategorized")
            entities[export_name] = {
                "entity": export_name,
                "category": cat,
                "total_fields": 0,
                "described_fields": 0,
                "fields": []
            }
        entities[export_name]["total_fields"] += 1
        if desc:
            entities[export_name]["described_fields"] += 1
        entities[export_name]["fields"].append({"field": field, "description": desc})
    
    # Category summary
    cats = defaultdict(lambda: {"entities": 0, "fields": 0, "described": 0, "entity_names": []})
    for name, info in sorted(entities.items()):
        c = info["category"]
        cats[c]["entities"] += 1
        cats[c]["fields"] += info["total_fields"]
        cats[c]["described"] += info["described_fields"]
        cats[c]["entity_names"].append(name)
    
    total_fields = sum(e["total_fields"] for e in entities.values())
    total_described = sum(e["described_fields"] for e in entities.values())
    
    output[sheet_name] = {
        "total_entities": len(entities),
        "total_fields": total_fields,
        "total_described": total_described,
        "description_pct": round(total_described / total_fields * 100, 1) if total_fields > 0 else 0,
        "categories": {k: {
            "entities": v["entities"],
            "fields": v["fields"],
            "described": v["described"],
            "entity_names": v["entity_names"]
        } for k, v in sorted(cats.items())},
        "entities": list(entities.values())
    }
    
    print(f"\n{'='*70}")
    print(f"  {sheet_name}")
    print(f"{'='*70}")
    print(f"  Total entities: {len(entities)}")
    print(f"  Total fields: {total_fields}")
    print(f"  Described: {total_described} ({output[sheet_name]['description_pct']}%)")
    print(f"\n  {'Category':<30} {'Entities':>8} {'Fields':>8} {'Described':>10}")
    print(f"  {'-'*60}")
    for cat_name in sorted(cats.keys()):
        s = cats[cat_name]
        pct = round(s['described']/s['fields']*100) if s['fields'] > 0 else 0
        print(f"  {cat_name:<30} {s['entities']:>8} {s['fields']:>8} {s['described']:>9} ({pct}%)")

wb.close()

with open("full-entity-inventory.json", "w") as f:
    json.dump(output, f, indent=2)

# Also generate markdown table for top entities
print("\n\n### Markdown table: Top entities by field count (Bulk Patient Export)")
bulk = output.get("Bulk Patient Export", output.get(list(output.keys())[-1]))
sorted_ents = sorted(bulk["entities"], key=lambda x: x["total_fields"], reverse=True)
print("| Entity | Fields | Described | Category |")
print("|---|---|---|---|")
for e in sorted_ents[:25]:
    print(f"| {e['entity']} | {e['total_fields']} | {e['described_fields']} | {e['category']} |")

print(f"\nFull inventory saved to full-entity-inventory.json")
