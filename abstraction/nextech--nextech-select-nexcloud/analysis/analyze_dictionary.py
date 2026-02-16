#!/usr/bin/env python3
"""
Parse the Nextech Select/NexCloud EHI Export Data Dictionary (XLSX)
and produce a full entity inventory JSON plus summary statistics.

Reads the XLSX directly (not the prior agent's extraction) to independently
verify field counts, description coverage, and content.
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

XLSX_PATH = Path(__file__).parent.parent.parent.parent / "results/nextech--nextech-select/downloads/Nextech_Select_NexCloud_EHI_Export_Data_Dictionary_2025.xlsx"
OUTPUT_DIR = Path(__file__).parent

def parse_xlsx():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    entities = []
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        
        # Find header row (first row with content)
        header = None
        header_idx = 0
        for i, row in enumerate(rows):
            non_empty = [c for c in row if c is not None and str(c).strip()]
            if len(non_empty) >= 2:
                header = [str(c).strip() if c else "" for c in row]
                header_idx = i
                break
        
        if header is None:
            continue
        
        # Determine export format from sheet name
        export_format = "Unknown"
        export_filename = None
        if "(CSV)" in sheet_name:
            export_format = "CSV"
            base = sheet_name.replace("(CSV)", "").strip()
            # Map sheet names to filenames
            name_map = {
                "Charges": "Charges.csv",
                "Appointments": "Appointments.csv",
                "Patient Demographics": "PatientDemographics.csv",
                "Payments": "Payments.csv",
                "Custom": "CustomData.csv",
                "Follow up": "Follow up.csv",
                "Insurances": "Insurances.csv",
                "Notes": "Notes.csv",
                "Recalls": "Recalls.csv",
                "Recalls Steps": "RecallsSteps.csv",
                "Payment Plans": "PaymentPlans.csv",
                "PracYakker": "PracYakker.csv",
                "iPad Tasks": "iPad Tasks.csv",
                "iPad Notes": "iPad Notes.csv",
            }
            export_filename = name_map.get(base, f"{base}.csv")
        elif "(PDF)" in sheet_name:
            export_format = "PDF"
            export_filename = "EMN_*.pdf"
        
        # Parse fields
        fields = []
        for row in rows[header_idx + 1:]:
            if row is None:
                continue
            values = [str(c).strip() if c is not None else "" for c in row]
            if not any(values):
                continue
            
            field = {"field_name": values[0] if len(values) > 0 else ""}
            if not field["field_name"]:
                continue
            
            # Map remaining columns based on header
            for j, h in enumerate(header):
                if j == 0:
                    continue
                h_lower = h.lower().strip()
                if j < len(values):
                    val = values[j]
                    if "description" in h_lower:
                        field["description"] = val
                    elif "note" in h_lower:
                        field["note"] = val if val else None
                    elif "type" in h_lower:
                        field["type"] = val if val else None
                    else:
                        field[h] = val if val else None
            
            # Ensure description key exists
            if "description" not in field:
                field["description"] = ""
            if "note" not in field:
                field["note"] = None
            
            fields.append(field)
        
        entity = {
            "sheet_name": sheet_name,
            "export_format": export_format,
            "export_filename": export_filename,
            "field_count": len(fields),
            "fields_with_description": sum(1 for f in fields if f.get("description", "").strip()),
            "fields_with_note": sum(1 for f in fields if f.get("note")),
            "fields": fields,
        }
        entities.append(entity)
    
    wb.close()
    return entities

def categorize_entity(sheet_name):
    """Assign a domain category based on entity name."""
    name = sheet_name.lower()
    if "charge" in name:
        return "Billing"
    elif "payment" in name:
        return "Billing"
    elif "insurance" in name:
        return "Insurance"
    elif "demographic" in name:
        return "Demographics"
    elif "appointment" in name:
        return "Scheduling"
    elif "emn" in name:
        return "Clinical Encounters"
    elif "note" in name:
        return "Clinical Notes"
    elif "recall" in name:
        return "Patient Outreach"
    elif "follow up" in name:
        return "Patient Outreach"
    elif "custom" in name:
        return "Custom Data"
    elif "pracyakker" in name:
        return "Internal Messaging"
    elif "dsi" in name:
        return "Clinical Decision Support"
    elif "ipad" in name:
        return "iPad Application"
    return "Other"

def main():
    entities = parse_xlsx()
    
    # Build full inventory
    inventory = {
        "source_file": XLSX_PATH.name,
        "extraction_date": "2026-02-16",
        "extraction_tool": "analyze_dictionary.py (independent extraction)",
        "total_entities": len(entities),
        "total_fields": sum(e["field_count"] for e in entities),
        "total_fields_with_description": sum(e["fields_with_description"] for e in entities),
        "total_fields_with_note": sum(e["fields_with_note"] for e in entities),
        "entities": []
    }
    
    for e in entities:
        cat = categorize_entity(e["sheet_name"])
        entry = {
            "sheet_name": e["sheet_name"],
            "export_format": e["export_format"],
            "export_filename": e["export_filename"],
            "category": cat,
            "field_count": e["field_count"],
            "fields_with_description": e["fields_with_description"],
            "fields_with_note": e["fields_with_note"],
            "description_coverage_pct": round(100 * e["fields_with_description"] / e["field_count"], 1) if e["field_count"] > 0 else 0,
            "fields": e["fields"]
        }
        inventory["entities"].append(entry)
    
    # Save full inventory
    with open(OUTPUT_DIR / "full-entity-inventory.json", "w") as f:
        json.dump(inventory, f, indent=2)
    
    # Summary statistics
    summary = {
        "total_entities": inventory["total_entities"],
        "total_fields": inventory["total_fields"],
        "total_fields_with_description": inventory["total_fields_with_description"],
        "description_coverage_pct": round(100 * inventory["total_fields_with_description"] / inventory["total_fields"], 1) if inventory["total_fields"] > 0 else 0,
        "total_fields_with_note": inventory["total_fields_with_note"],
        "entities_by_category": {},
        "entity_summary": []
    }
    
    # Category breakdown
    cats = {}
    for e in inventory["entities"]:
        cat = e["category"]
        if cat not in cats:
            cats[cat] = {"entity_count": 0, "field_count": 0, "fields_with_description": 0}
        cats[cat]["entity_count"] += 1
        cats[cat]["field_count"] += e["field_count"]
        cats[cat]["fields_with_description"] += e["fields_with_description"]
    summary["entities_by_category"] = cats
    
    # Per-entity summary (no field details)
    for e in inventory["entities"]:
        summary["entity_summary"].append({
            "sheet_name": e["sheet_name"],
            "export_format": e["export_format"],
            "export_filename": e["export_filename"],
            "category": e["category"],
            "field_count": e["field_count"],
            "fields_with_description": e["fields_with_description"],
            "description_coverage_pct": e["description_coverage_pct"],
        })
    
    with open(OUTPUT_DIR / "summary-statistics.json", "w") as f:
        json.dump(summary, f, indent=2)
    
    # Print summary to stdout
    print(f"Total entities: {summary['total_entities']}")
    print(f"Total fields: {summary['total_fields']}")
    print(f"Fields with description: {summary['total_fields_with_description']} ({summary['description_coverage_pct']}%)")
    print(f"Fields with note: {summary['total_fields_with_note']}")
    print()
    print("Category breakdown:")
    for cat, info in sorted(cats.items()):
        print(f"  {cat}: {info['entity_count']} entities, {info['field_count']} fields, {info['fields_with_description']} described")
    print()
    print("Per-entity:")
    for e in summary["entity_summary"]:
        print(f"  {e['sheet_name']}: {e['field_count']} fields ({e['fields_with_description']} described, {e['description_coverage_pct']}%)")

if __name__ == "__main__":
    main()
