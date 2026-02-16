#!/usr/bin/env python3
"""
Parse Nextech Select/NexCloud EHI Export Data Dictionary XLSX into
entity-inventory-full.json and entity-inventory-summary.json.
"""

import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

XLSX_PATH = Path(__file__).parent.parent / "downloads" / "Nextech_Select_NexCloud_EHI_Export_Data_Dictionary_2025.xlsx"
OUT_DIR = Path(__file__).parent

def parse_sheet(ws):
    """Parse a single worksheet into a list of field dicts."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    
    # Find header row
    headers = [str(h).strip() if h else "" for h in rows[0]]
    
    fields = []
    for row in rows[1:]:
        if not any(row):
            continue
        field = {}
        for i, h in enumerate(headers):
            val = row[i] if i < len(row) else None
            if val is not None:
                val = str(val).strip()
                if val == "":
                    val = None
            field[h] = val
        
        # Normalize to standard keys
        normalized = {
            "field_name": field.get("Field Name") or field.get("Field name") or field.get("field_name") or "",
            "description": field.get("Description") or field.get("description") or "",
            "note": field.get("Note") or field.get("note") or None,
        }
        # Include any extra columns
        for k, v in field.items():
            if k and k not in ("Field Name", "Field name", "field_name", "Description", "description", "Note", "note", ""):
                normalized[k.lower().replace(" ", "_")] = v
        
        if normalized["field_name"]:
            fields.append(normalized)
    
    return fields

def categorize_entity(sheet_name: str) -> str:
    """Assign a domain category based on sheet/entity name."""
    name_lower = sheet_name.lower()
    if "charge" in name_lower:
        return "Billing"
    elif "payment" in name_lower:
        return "Billing"
    elif "insurance" in name_lower:
        return "Insurance"
    elif "demograph" in name_lower:
        return "Demographics"
    elif "appointment" in name_lower:
        return "Scheduling"
    elif "emn" in name_lower:
        return "Clinical Encounters"
    elif "note" in name_lower:
        return "Clinical Notes"
    elif "recall" in name_lower:
        return "Patient Outreach"
    elif "follow" in name_lower:
        return "Patient Outreach"
    elif "custom" in name_lower:
        return "Custom Data"
    elif "pracyakker" in name_lower:
        return "Internal Messaging"
    elif "ipad" in name_lower:
        return "iPad Application"
    elif "ccda" in name_lower:
        return "Clinical (CCDA)"
    elif "dsi" in name_lower or "feedback" in name_lower:
        return "Clinical Decision Support"
    else:
        return "Other"

def main():
    wb = openpyxl.load_workbook(XLSX_PATH, read_only=True, data_only=True)
    
    entities = []
    total_fields = 0
    total_with_desc = 0
    total_with_note = 0
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        fields = parse_sheet(ws)
        
        # Parse format from sheet name
        import re
        fmt_match = re.search(r'\((\w+)\)\s*$', sheet_name)
        export_format = fmt_match.group(1) if fmt_match else "Unknown"
        base_name = re.sub(r'\s*\(\w+\)\s*$', '', sheet_name).strip()
        ext = export_format.lower()
        export_filename = f"{base_name}.{ext}" if ext != "unknown" else base_name
        
        desc_count = sum(1 for f in fields if f.get("description"))
        note_count = sum(1 for f in fields if f.get("note"))
        
        entity = {
            "entity_name": base_name,
            "sheet_name": sheet_name,
            "export_format": export_format,
            "export_filename": export_filename,
            "category": categorize_entity(sheet_name),
            "field_count": len(fields),
            "fields_with_description": desc_count,
            "fields_with_note": note_count,
            "fields": fields
        }
        entities.append(entity)
        total_fields += len(fields)
        total_with_desc += desc_count
        total_with_note += note_count
    
    wb.close()
    
    # Full inventory
    full_inventory = {
        "source_file": "Nextech_Select_NexCloud_EHI_Export_Data_Dictionary_2025.xlsx",
        "extraction_date": "2026-02-16",
        "total_entities": len(entities),
        "total_fields": total_fields,
        "total_fields_with_description": total_with_desc,
        "total_fields_with_note": total_with_note,
        "entities": entities
    }
    
    with open(OUT_DIR / "entity-inventory-full.json", "w") as f:
        json.dump(full_inventory, f, indent=2)
    
    # Summary
    category_breakdown = {}
    for e in entities:
        cat = e["category"]
        if cat not in category_breakdown:
            category_breakdown[cat] = {"entity_count": 0, "field_count": 0, "fields_with_description": 0}
        category_breakdown[cat]["entity_count"] += 1
        category_breakdown[cat]["field_count"] += e["field_count"]
        category_breakdown[cat]["fields_with_description"] += e["fields_with_description"]
    
    entity_summary = []
    for e in entities:
        entity_summary.append({
            "entity_name": e["entity_name"],
            "sheet_name": e["sheet_name"],
            "export_format": e["export_format"],
            "export_filename": e["export_filename"],
            "category": e["category"],
            "field_count": e["field_count"],
            "fields_with_description": e["fields_with_description"],
            "fields_with_note": e["fields_with_note"],
        })
    
    summary = {
        "source_file": full_inventory["source_file"],
        "extraction_date": full_inventory["extraction_date"],
        "total_entities": len(entities),
        "total_fields": total_fields,
        "total_fields_with_description": total_with_desc,
        "total_fields_with_note": total_with_note,
        "description_coverage_pct": round(total_with_desc / total_fields * 100, 1) if total_fields > 0 else 0,
        "category_breakdown": category_breakdown,
        "entities": entity_summary
    }
    
    with open(OUT_DIR / "entity-inventory-summary.json", "w") as f:
        json.dump(summary, f, indent=2)
    
    # Print summary
    print(f"Total entities: {len(entities)}")
    print(f"Total fields: {total_fields}")
    print(f"Fields with description: {total_with_desc} ({round(total_with_desc/total_fields*100,1)}%)")
    print(f"Fields with note: {total_with_note} ({round(total_with_note/total_fields*100,1)}%)")
    print()
    print("Entities:")
    for e in entity_summary:
        print(f"  {e['entity_name']:30s} {e['export_format']:5s} {e['field_count']:4d} fields  ({e['category']})")
    print()
    print("Category breakdown:")
    for cat, info in sorted(category_breakdown.items()):
        print(f"  {cat:30s} {info['entity_count']:2d} entities, {info['field_count']:4d} fields")

if __name__ == "__main__":
    main()
