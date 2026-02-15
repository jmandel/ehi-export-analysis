#!/usr/bin/env python3
"""Parse the Excel data dictionary to count entities, fields, and description quality.
Run with: .venv/bin/python3 parse_xlsx.py"""

import json
import sys
from pathlib import Path
import openpyxl

XLSX_PATH = Path("/home/jmandel/hobby/ehi-export-analysis/results/isalus-healthcare/downloads/data-elements/iSalus_EHI_ExportDataElements_published_version1_Oct2023_pristine.xlsx")

wb = openpyxl.load_workbook(XLSX_PATH, read_only=True)

print(f"=== XLSX Data Dictionary ===")
print(f"Sheets: {wb.sheetnames}")
print()

all_entities = {}
total_fields = 0
fields_with_desc = 0

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        continue

    # Print header
    header = rows[0]
    print(f"Sheet: {sheet_name}")
    print(f"  Header: {header}")
    print(f"  Data rows: {len(rows) - 1}")

    # Find column indices
    header_lower = [str(h).lower().strip() if h else "" for h in header]
    export_name_col = None
    field_col = None
    desc_col = None

    for i, h in enumerate(header_lower):
        if "export" in h and "name" in h:
            export_name_col = i
        elif h == "field":
            field_col = i
        elif "description" in h:
            desc_col = i

    if export_name_col is None or field_col is None:
        print(f"  WARNING: Could not find expected columns")
        continue

    for row in rows[1:]:
        if row[export_name_col] is None:
            continue
        entity = str(row[export_name_col]).strip()
        field = str(row[field_col]).strip() if row[field_col] else ""
        desc = str(row[desc_col]).strip() if desc_col is not None and row[desc_col] else ""

        if entity not in all_entities:
            all_entities[entity] = {"fields": [], "described": 0, "total": 0}

        all_entities[entity]["total"] += 1
        total_fields += 1
        if desc and desc.lower() != "none":
            all_entities[entity]["described"] += 1
            fields_with_desc += 1
        all_entities[entity]["fields"].append({"field": field, "description": desc})

print()
print(f"=== Summary ===")
print(f"Total entities: {len(all_entities)}")
print(f"Total fields: {total_fields}")
print(f"Fields with descriptions: {fields_with_desc}")
print(f"Fields without descriptions: {total_fields - fields_with_desc}")
print()

print(f"{'Entity':<50} {'Total':>6} {'Described':>10}")
print("-" * 70)
for entity, info in sorted(all_entities.items()):
    print(f"{entity:<50} {info['total']:>6} {info['described']:>10}")

# Save structured output
output_path = Path("/home/jmandel/hobby/ehi-export-analysis/abstraction/isalus-healthcare/analysis/xlsx_summary.json")
summary = {
    "total_entities": len(all_entities),
    "total_fields": total_fields,
    "fields_with_descriptions": fields_with_desc,
    "entities": {k: {"total_fields": v["total"], "described_fields": v["described"]} for k, v in sorted(all_entities.items())}
}
with open(output_path, "w") as fh:
    json.dump(summary, fh, indent=2)
print(f"\nJSON saved to {output_path}")
