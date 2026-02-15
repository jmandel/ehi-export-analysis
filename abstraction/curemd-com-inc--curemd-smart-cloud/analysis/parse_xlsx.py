#!/usr/bin/env python3
"""Parse CureMD EHI_Export_DRS.xlsx independently to verify field counts and descriptions."""
import json
import sys

try:
    from openpyxl import load_workbook
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    from openpyxl import load_workbook

XLSX_PATH = "/home/jmandel/hobby/ehi-export-analysis/results/curemd-com-inc--curemd-smart-cloud/downloads/EHI_Export_DRS.xlsx"

wb = load_workbook(XLSX_PATH, read_only=True)
print("Sheets:", wb.sheetnames)

# Parse EHI Export sheet
ehi_sheet = wb["EHI Export"]
ehi_rows = list(ehi_sheet.iter_rows(values_only=True))
print(f"\nEHI Export sheet: {len(ehi_rows)} rows")
for r in ehi_rows:
    vals = [str(c) for c in r if c is not None]
    if vals:
        print("  ", " | ".join(vals[:3]))

# Parse DRS sheet
drs_sheet = wb["DRS"]
drs_rows = list(drs_sheet.iter_rows(values_only=True))
print(f"\nDRS sheet: {len(drs_rows)} rows (including header)")

# Header
header = drs_rows[0]
print(f"Header: {header}")

# Parse data classes
data_classes = {}
current_class = None
total_fields = 0
fields_with_desc = 0
fields_with_type = 0

for row in drs_rows[1:]:
    class_name = row[0]
    field_name = row[1]
    data_type = row[2]
    description = row[3] if len(row) > 3 else None

    if class_name is not None and str(class_name).strip():
        current_class = str(class_name).strip()
        if current_class not in data_classes:
            data_classes[current_class] = {"fields": [], "fields_with_desc": 0, "fields_with_type": 0}

    if field_name is not None and str(field_name).strip() and str(field_name).strip() != "-":
        total_fields += 1
        field_info = {
            "name": str(field_name).strip(),
            "type": str(data_type).strip() if data_type else None,
            "description": str(description).strip() if description else None
        }
        if current_class:
            data_classes[current_class]["fields"].append(field_info)
            if description and str(description).strip():
                data_classes[current_class]["fields_with_desc"] += 1
                fields_with_desc += 1
            if data_type and str(data_type).strip():
                data_classes[current_class]["fields_with_type"] += 1
                fields_with_type += 1

# Also count special classes (Documents & Images, Provider Notes) which have "-" as field
for row in drs_rows[1:]:
    class_name = row[0]
    field_name = row[1]
    if class_name and str(class_name).strip() and (field_name is None or str(field_name).strip() == "-"):
        cn = str(class_name).strip()
        if cn not in data_classes:
            data_classes[cn] = {"fields": [], "fields_with_desc": 0, "fields_with_type": 0, "special": True}

print(f"\nTotal data classes: {len(data_classes)}")
print(f"Total fields (excluding special classes): {total_fields}")
print(f"Fields with descriptions: {fields_with_desc} ({fields_with_desc/total_fields*100:.1f}%)")
print(f"Fields with data types: {fields_with_type} ({fields_with_type/total_fields*100:.1f}%)")

print("\n--- Data Class Summary ---")
summary = []
for cls_name, cls_data in data_classes.items():
    fc = len(cls_data["fields"])
    fd = cls_data["fields_with_desc"]
    ft = cls_data["fields_with_type"]
    special = cls_data.get("special", False)
    print(f"  {cls_name}: {fc} fields, {fd} described, {ft} typed" + (" [SPECIAL]" if special else ""))
    summary.append({
        "name": cls_name,
        "field_count": fc,
        "fields_with_descriptions": fd,
        "fields_with_types": ft,
        "special": special,
        "fields": cls_data["fields"]
    })

# Check unique data types
all_types = set()
for cls_data in data_classes.values():
    for f in cls_data["fields"]:
        if f["type"]:
            all_types.add(f["type"])
print(f"\nUnique data types: {sorted(all_types)}")

# Save full inventory
output = {
    "total_data_classes": len(data_classes),
    "total_fields": total_fields,
    "fields_with_descriptions": fields_with_desc,
    "fields_with_types": fields_with_type,
    "pct_described": round(fields_with_desc/total_fields*100, 1),
    "unique_data_types": sorted(all_types),
    "data_classes": summary
}

with open("full-entity-inventory.json", "w") as f:
    json.dump(output, f, indent=2)

print("\nSaved full-entity-inventory.json")
wb.close()
