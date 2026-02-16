#!/usr/bin/env python3
"""Parse all three Nextech EHI Export XLSX data dictionaries into a unified JSON inventory."""

import json
import openpyxl
from pathlib import Path

DOWNLOADS = Path(__file__).resolve().parent.parent.parent.parent / "results" / "nextech--nextech-ehr-icp" / "downloads"
OUTPUT = Path(__file__).resolve().parent

PLATFORMS = [
    {
        "platform": "IntelleChartPRO (ICP)",
        "certified_product": "Nextech EHR (ICP)",
        "chpl_id": 11724,
        "file": "nextech-ehr-icp-ehi-data-dictionary.xlsx",
    },
    {
        "platform": "SRSPro",
        "certified_product": "SRS EHR",
        "chpl_id": 11040,
        "file": "srspro-ehi-data-dictionary.xlsx",
    },
    {
        "platform": "Nextech Select/NexCloud",
        "certified_product": None,  # Not separately CHPL-certified
        "chpl_id": None,
        "file": "nextech-select-nexcloud-ehi-data-dictionary.xlsx",
    },
]


def parse_sheet(ws):
    """Parse a single worksheet into a list of field dicts."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], None

    # Find the header row (first row with at least 2 non-None values)
    header_idx = None
    for i, row in enumerate(rows):
        non_none = [c for c in row if c is not None]
        if len(non_none) >= 2:
            header_idx = i
            break
    if header_idx is None:
        return [], None

    headers = [str(c).strip() if c else f"col_{j}" for j, c in enumerate(rows[header_idx])]

    # Detect export format from header or first few rows
    export_format = None
    for row in rows[:header_idx + 1]:
        for cell in row:
            if cell and isinstance(cell, str):
                cl = cell.lower()
                if 'json' in cl:
                    export_format = 'JSON'
                elif 'csv' in cl:
                    export_format = 'CSV'
                elif 'xml' in cl:
                    export_format = 'XML'
                elif 'pdf' in cl:
                    export_format = 'PDF'
                elif 'txt' in cl:
                    export_format = 'TXT'

    fields = []
    for row in rows[header_idx + 1:]:
        if all(c is None for c in row):
            continue
        field = {}
        for j, val in enumerate(row):
            if j < len(headers):
                key = headers[j]
                field[key] = str(val).strip() if val is not None else None
        # Only include if there's at least a field name
        name_candidates = [field.get(h) for h in headers if 'name' in h.lower() or 'field' in h.lower() or 'column' in h.lower()]
        if not name_candidates:
            name_candidates = [field.get(headers[0])] if headers else []
        if any(n and n != 'None' for n in name_candidates):
            fields.append(field)

    return fields, export_format


def parse_workbook(filepath):
    """Parse an entire workbook into entities."""
    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    entities = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        fields, fmt = parse_sheet(ws)
        # Determine description column
        desc_count = 0
        for f in fields:
            has_desc = False
            for k, v in f.items():
                if v and v != 'None' and any(d in k.lower() for d in ['description', 'desc', 'note', 'comment', 'definition']):
                    has_desc = True
                    break
            if has_desc:
                desc_count += 1

        entities.append({
            "entity_name": sheet_name.strip(),
            "field_count": len(fields),
            "fields_with_descriptions": desc_count,
            "export_format": fmt,
            "fields": fields,
        })
    wb.close()
    return entities


def main():
    inventory = {
        "extraction_date": "2026-02-16",
        "platforms": [],
        "summary": {},
    }

    total_entities = 0
    total_fields = 0
    total_described = 0

    for pinfo in PLATFORMS:
        fpath = DOWNLOADS / pinfo["file"]
        entities = parse_workbook(fpath)
        plat_fields = sum(e["field_count"] for e in entities)
        plat_described = sum(e["fields_with_descriptions"] for e in entities)

        platform_data = {
            "platform": pinfo["platform"],
            "certified_product": pinfo["certified_product"],
            "chpl_id": pinfo["chpl_id"],
            "source_file": pinfo["file"],
            "entity_count": len(entities),
            "total_fields": plat_fields,
            "fields_with_descriptions": plat_described,
            "entities": entities,
        }
        inventory["platforms"].append(platform_data)

        total_entities += len(entities)
        total_fields += plat_fields
        total_described += plat_described

        print(f"\n{pinfo['platform']}:")
        print(f"  Entities: {len(entities)}, Fields: {plat_fields}, Described: {plat_described}")
        for e in entities:
            print(f"    {e['entity_name']}: {e['field_count']} fields ({e['fields_with_descriptions']} described), format={e['export_format']}")

    inventory["summary"] = {
        "total_platforms": len(PLATFORMS),
        "total_entities": total_entities,
        "total_fields": total_fields,
        "total_fields_with_descriptions": total_described,
        "description_rate": f"{total_described/total_fields*100:.1f}%" if total_fields > 0 else "N/A",
    }

    print(f"\n=== OVERALL ===")
    print(f"Total entities: {total_entities}")
    print(f"Total fields: {total_fields}")
    print(f"Description rate: {inventory['summary']['description_rate']}")

    with open(OUTPUT / "full-entity-inventory.json", "w") as f:
        json.dump(inventory, f, indent=2)
    print(f"\nSaved to {OUTPUT / 'full-entity-inventory.json'}")


if __name__ == "__main__":
    main()
